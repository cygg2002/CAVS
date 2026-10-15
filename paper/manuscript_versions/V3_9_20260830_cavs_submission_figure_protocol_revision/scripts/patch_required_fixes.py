#!/usr/bin/env python3
"""Apply the required post-appendix-removal fixes to the current manuscript DOCX."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "MANUSCRIPT_V3_9_ZH_FINAL_S6_S7_VOC_BASELINES_FIG5_E2E_LABELS_FIXED.docx"
OUTPUT = ROOT / "MANUSCRIPT_V3_9_ZH_FINAL_REQUIRED_FIXES_REVISED.docx"


def math_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("m:t")))


def math_run(text: str, *, normal: bool = False):
    run = OxmlElement("m:r")
    if normal:
        properties = OxmlElement("m:rPr")
        properties.append(OxmlElement("m:nor"))
        run.append(properties)
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def math_subscript(base: str, subscript: str):
    scripted = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e")
    base_node.append(math_run(base))
    sub_node = OxmlElement("m:sub")
    sub_node.append(math_run(subscript))
    scripted.extend((base_node, sub_node))
    return scripted


def math_subsup(base_element, subscript: str, superscript: str):
    scripted = OxmlElement("m:sSubSup")
    base_node = OxmlElement("m:e")
    base_node.append(base_element)
    sub_node = OxmlElement("m:sub")
    sub_node.append(math_run(subscript))
    sup_node = OxmlElement("m:sup")
    sup_node.append(math_run(superscript))
    scripted.extend((base_node, sub_node, sup_node))
    return scripted


def math_delimiter(begin: str, end: str, *contents):
    delimiter = OxmlElement("m:d")
    properties = OxmlElement("m:dPr")
    begin_node = OxmlElement("m:begChr")
    begin_node.set(qn("m:val"), begin)
    end_node = OxmlElement("m:endChr")
    end_node.set(qn("m:val"), end)
    properties.extend((begin_node, end_node))
    body = OxmlElement("m:e")
    body.extend(contents)
    delimiter.extend((properties, body))
    return delimiter


def inline_math(*contents):
    equation = OxmlElement("m:oMath")
    equation.extend(contents)
    return equation


def formula_dataset():
    image = math_subscript("I", "i")
    image_set = math_delimiter("{", "}", image)
    indexed_set = math_subsup(image_set, "i=1", "N")
    return inline_math(math_run("D="), indexed_set)


def formula_cloud_action():
    return inline_math(
        math_run("a∈"),
        math_subscript("A", "i"),
        math_run("∖"),
        math_delimiter("{", "}", math_run("e")),
    )


def formula_k_limit():
    return inline_math(
        math_run("K="),
        math_delimiter("⌊", "⌋", math_run("B/c")),
    )


def formula_full_complexity():
    return inline_math(
        math_run("O(TN"),
        math_delimiter("|", "|", math_run("A")),
        math_run("+U"),
        math_run("log", normal=True),
        math_run("U)"),
    )


def formula_online_complexity():
    return inline_math(
        math_run("O("),
        math_delimiter("|", "|", math_run("A")),
        math_run(")"),
    )


def formula_context_edge_results():
    context = math_subscript("C", "i")
    scripted = OxmlElement("m:sSubSup")
    base_node = OxmlElement("m:e")
    base_node.append(math_run("F"))
    sub_node = OxmlElement("m:sub")
    sub_node.append(context)
    sup_node = OxmlElement("m:sup")
    sup_node.append(math_run("e"))
    scripted.extend((base_node, sub_node, sup_node))
    return inline_math(scripted)


def replace_math(paragraph, old_text: str, replacement) -> None:
    matches = [
        node
        for node in paragraph._p.iterchildren()
        if node.tag == qn("m:oMath") and math_text(node) == old_text
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one inline equation {old_text!r}, found {len(matches)} in {paragraph.text!r}"
        )
    matches[0].getparent().replace(matches[0], replacement)


def set_text_node(node, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set(qn("xml:space"), "preserve")
    elif qn("xml:space") in node.attrib:
        del node.attrib[qn("xml:space")]


def normal_text(paragraph) -> str:
    return "".join(node.text or "" for node in paragraph._p.iter(qn("w:t")))


def replace_normal_text(paragraph, old: str, new: str) -> None:
    nodes = list(paragraph._p.iter(qn("w:t")))
    values = [node.text or "" for node in nodes]
    joined = "".join(values)
    if joined.count(old) != 1:
        raise RuntimeError(
            f"Expected one occurrence of {old!r}, found {joined.count(old)} in {joined!r}"
        )
    start = joined.index(old)
    end = start + len(old)
    cursor = 0
    first_index = last_index = None
    first_offset = last_offset = 0
    for index, value in enumerate(values):
        next_cursor = cursor + len(value)
        if first_index is None and start < next_cursor:
            first_index = index
            first_offset = start - cursor
        if end <= next_cursor:
            last_index = index
            last_offset = end - cursor
            break
        cursor = next_cursor
    if first_index is None or last_index is None:
        raise RuntimeError(f"Could not map replacement range for {old!r}")
    if first_index == last_index:
        value = values[first_index]
        set_text_node(
            nodes[first_index], value[:first_offset] + new + value[last_offset:]
        )
        return
    first_value = values[first_index]
    last_value = values[last_index]
    set_text_node(nodes[first_index], first_value[:first_offset] + new)
    for index in range(first_index + 1, last_index):
        set_text_node(nodes[index], "")
    set_text_node(nodes[last_index], last_value[last_offset:])


def clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def append_plain_run(paragraph, text: str):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "STSong")
    fonts.set(qn("w:hint"), "eastAsia")
    return run


def find_paragraph(document, prefix: str):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def find_paragraph_with_math(document, equation_text: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if any(
            node.tag == qn("m:oMath") and math_text(node) == equation_text
            for node in paragraph._p.iterchildren()
        )
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph containing equation {equation_text!r}, found {len(matches)}"
        )
    return matches[0]


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(OUTPUT)

    document = Document(SOURCE)

    # Repair the three formula fragments that produced red compatibility errors
    # in the independent LibreOffice render.
    replace_math(find_paragraph(document, "给定输入序列或批次"), "D={Ii}i=1N", formula_dataset())
    replace_math(find_paragraph(document, "对每个可行云动作"), "a∈Ai\\{e}", formula_cloud_action())
    complexity = find_paragraph(document, "集合编码至多处理 64 个框")
    replace_math(complexity, "O(TN|A|+UlogU)", formula_full_complexity())
    replace_math(complexity, "O(|A|)", formula_online_complexity())

    # Define K in the equal-cost special case and remove the deleted appendix reference.
    special_case = find_paragraph(document, "因此，10% 指 Always-Original")
    clear_paragraph_content(special_case)
    append_plain_run(
        special_case,
        "因此，10% 指 Always-Original JPEG payload 总量的 10%，而非上传 10% 的图像。"
        "若每幅图像仅有一个成本均为 ",
    )
    special_case._p.append(inline_math(math_run("c")))
    append_plain_run(special_case, " 的云动作，则式（3）退化为选择至多 ")
    special_case._p.append(formula_k_limit())
    append_plain_run(special_case, " 个正价值样本；若正价值样本少于 ")
    special_case._p.append(inline_math(math_run("K")))
    append_plain_run(
        special_case,
        "，其余图像保留 Edge。一旦候选动作成本不同，成员与动作通常需要联合求解。",
    )

    # Restore the minimum architecture and optimization details needed to
    # reproduce the value model after the appendix was removed.
    representation = find_paragraph(document, "CAVS 由检测集合编码器")
    replace_normal_text(
        representation,
        "类别使用可学习嵌入。",
        "类别使用可学习嵌入。正式网络的类别嵌入为 8 维；逐框编码器和动作编码器均为两层 MLP、"
        "隐藏维度为 32，图像—动作联合表示为 64 维，并采用 SiLU 激活和 LayerNorm，输出头的 dropout 为 0.05。",
    )

    reward = find_paragraph(document, "ORIC 在代表性图像上下文中定义")
    final_run = reward._p[-1]
    if final_run.tag != qn("w:r") or normal_text(reward).endswith("，对云动作定义原始上下文奖励") is False:
        raise RuntimeError("Unexpected reward paragraph structure")
    final_text_nodes = list(final_run.iter(qn("w:t")))
    if len(final_text_nodes) != 1 or final_text_nodes[0].text != "，对云动作定义原始上下文奖励":
        raise RuntimeError("Unexpected final reward run")
    set_text_node(final_text_nodes[0], "，记 ")
    reward._p.append(formula_context_edge_results())
    following_run = deepcopy(final_run)
    following_text = list(following_run.iter(qn("w:t")))
    set_text_node(
        following_text[0],
        " 为上下文图像在 Edge 动作下形成的检测结果集合，并对云动作定义原始上下文奖励",
    )
    reward._p.append(following_run)

    objective = find_paragraph_with_math(document, "siab=sgn(yia-yib)")
    replace_normal_text(
        objective,
        "其中 ",
        "其中，Ω 表示回归项中的全部可行图像—动作对（包括 Edge）；Π 表示同一图像内按预先固定动作顺序"
        "每对只计一次、且标准化目标差绝对值大于 10⁻⁸ 的可行动作对。",
    )
    replace_normal_text(
        objective,
        "，正式配置取 ",
        "。若 Π 为空或其权重和为 0，则将排序项定义为 0；正式配置取 ",
    )

    training = find_paragraph(document, "各 CAVS 配置均使用五折折外预测")
    replace_normal_text(
        training,
        "batch size 192，最多训练",
        "batch size 192，梯度范数裁剪阈值为 5，最多训练",
    )

    # Remove all surviving claims that depended on the deleted network-stress
    # section, while retaining the ordinary online testbed protocol and limits.
    testbed = find_paragraph(document, "所有在线方法使用相同的 Orange Pi")
    replace_normal_text(
        testbed,
        "网络敏感性实验仅控制往返时延和可用带宽，同时保持无丢包、无云端失败和无截止时间回退，"
        "且调度器不读取网络状态。因此该实验衡量传输条件对时延与吞吐的影响，而不评价网络状态自适应调度。",
        "",
    )

    discussion = find_paragraph(document, "CAVS 还承担 mandatory weak-detector")
    replace_normal_text(discussion, "，在 50 ms/5 Mbps 下 25% 档达到 350.02 ms", "")
    replace_normal_text(discussion, "这些结果表明", "这一结果表明")

    conclusion = find_paragraph(document, "结果也界定了适用范围")
    replace_normal_text(
        conclusion,
        "在无丢包、无云失败、无 deadline fallback 且动作固定的网络压力实验中，链路恶化只改变时延和吞吐，"
        "这一结论不能外推到会触发回退的弱网。",
        "",
    )

    # Add missing body references and the one uncited software reference.
    hardware = find_paragraph(document, "板端使用 Orange Pi 5 Ultra")
    replace_normal_text(hardware, "PyTorch 2.1.2", "PyTorch 2.1.2 [32]")

    comparison = find_paragraph(document, "比较方法包括 Edge-only")
    replace_normal_text(
        comparison,
        "为匹配统一字节预算，其适配版本按图像分数选择 Original，原生阈值工作点则单独报告。",
        "为匹配统一字节预算，其适配版本按图像分数选择 Original；原生阈值工作点不纳入固定字节主比较，"
        "VOC 与在线测试床中的相应结果仅作为预算外参考。",
    )

    main_results = find_paragraph(document, "图 2 汇总三组端—云检测器")
    replace_normal_text(
        main_results,
        "图 2 汇总三组端—云检测器在冻结协议下的点估计结果。",
        "图 2 汇总三组端—云检测器在冻结协议下的点估计结果，完整点结果见表 1。",
    )

    profile = find_paragraph(document, "为评估 ROI 几何规则与融合方式")
    replace_normal_text(
        profile,
        "我们在同一 COCO 划分、云模型、fold、种子和调度器下",
        "结果见表 5。我们在同一 COCO 划分、云模型、fold、种子和调度器下",
    )

    cache = find_paragraph(document, "全信息缓存参考为")
    replace_normal_text(cache, "全信息缓存参考为", "如图 4 所示，全信息缓存参考为")

    online = find_paragraph(document, "图 5 综合展示")
    replace_normal_text(online, "图 5 综合展示", "图 5 和表 7 综合展示")

    # Assertions make accidental scope expansion visible.
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    forbidden = (
        "附录 A.1",
        "网络敏感性实验",
        "在 50 ms/5 Mbps 下 25% 档达到 350.02 ms",
        "网络压力实验",
    )
    for phrase in forbidden:
        if phrase in all_text:
            raise RuntimeError(f"Stale phrase remains: {phrase}")
    required = (
        "梯度范数裁剪阈值为 5",
        "完整点结果见表 1",
        "结果见表 5",
        "图 5 和表 7",
        "PyTorch 2.1.2 [32]",
        "Statements and Declarations",
    )
    for phrase in required[:-1]:
        if phrase not in all_text:
            raise RuntimeError(f"Required phrase is missing: {phrase}")
    if required[-1] in all_text:
        raise RuntimeError("Statements and Declarations must remain absent")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
