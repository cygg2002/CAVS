#!/usr/bin/env python3
"""Build the V3.3 Chinese manuscript as an auditable academic DOCX artifact."""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "MANUSCRIPT_V3_3_ZH_DRAFT.md"
OUTPUT = ROOT / "MANUSCRIPT_V3_3_ZH_DRAFT.docx"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 72, "bottom": 72, "start": 120, "end": 120}

# Chinese-journal author-manuscript palette: all printable text is black.
BLACK = "000000"
RULE_GRAY = "7F7F7F"
CALLOUT_FILL = "F2F2F2"
PENDING_FILL = "F2F2F2"
PENDING_TEXT = BLACK

# The official Chinese template uses Song/Hei families and Times New Roman.
# Songti SC and Heiti SC are the installed macOS equivalents used for verified
# rendering; Latin glyphs and numerals remain Times New Roman.
BODY_FONT = "Times New Roman"
EAST_ASIA_FONT = "Songti SC"
HEADING_EAST_ASIA_FONT = "Heiti SC"


def set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    east_asia: str = EAST_ASIA_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:hint"), "eastAsia")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(
    style,
    *,
    size: float,
    color: str = BLACK,
    bold: bool = False,
    name: str = BODY_FONT,
    east_asia: str = EAST_ASIA_FONT,
    italic: bool = False,
) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:hint"), "eastAsia")


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, *, side: str, val: str, color: str = BLACK, size: int = 0) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    edge = borders.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        borders.append(edge)
    edge.set(qn("w:val"), val)
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "0")
    edge.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, *, side: str, color: str, size: int = 8) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    edge = pbdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        pbdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "6")
    edge.set(qn("w:color"), color)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, display, fld_end])
    set_run_font(run, size=8.5, color=BLACK)


def preprocess_markdown(source: str) -> str:
    """Prepare a temporary Word-safe source without changing the canonical Markdown."""
    text = source.replace(".svg)", ".png)")
    # Pandoc otherwise creates a figure caption from alt text and duplicates our explicit caption.
    text = re.sub(
        r"!\[[^\]]*\]\((figures/[^)]+\.png)\)",
        r"![](\1){width=6.15in}",
        text,
    )
    # Pandoc's DOCX writer drops TeX \tag; keep the manuscript's equation numbers visible.
    text = re.sub(r"\\tag\{([^}]+)\}", r"\\qquad\\text{(\1)}", text)
    # Remove only immediately repeated headings (one accidental duplicate exists in the draft).
    text = re.sub(r"(?m)^(#{1,6} .+)\n\1\n", r"\1\n", text)
    return text


def run_pandoc(temp_markdown: Path, raw_docx: Path) -> None:
    pandoc = os.environ.get("PANDOC") or shutil.which("pandoc") or str(Path.home() / ".local/bin/pandoc")
    command = [
        pandoc,
        str(temp_markdown),
        "--from=markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+raw_attribute",
        "--to=docx",
        "--standalone",
        f"--resource-path={ROOT}",
        "--dpi=300",
        "--output",
        str(raw_docx),
    ]
    subprocess.run(command, check=True, cwd=ROOT)


def configure_page(section) -> None:
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = Inches(1.0236)
    section.right_margin = Inches(0.7874)
    section.bottom_margin = Inches(1.0236)
    section.left_margin = Inches(0.7874)
    section.header_distance = Inches(0.63)
    section.footer_distance = Inches(0.63)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.first_line_indent = Pt(21)

    title = styles["Title"]
    set_style_font(
        title,
        size=22,
        color=BLACK,
        bold=False,
        east_asia=HEADING_EAST_ASIA_FONT,
    )
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.keep_with_next = True

    heading_tokens = {
        "Heading 1": (14, HEADING_EAST_ASIA_FONT, True, 12, 6),
        "Heading 2": (10.5, HEADING_EAST_ASIA_FONT, True, 8, 4),
        "Heading 3": (10.5, EAST_ASIA_FONT, True, 6, 3),
    }
    for name, (size, east_asia, bold, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style, size=size, color=BLACK, bold=bold, east_asia=east_asia)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.first_line_indent = Pt(0)

    caption = styles["Caption"]
    set_style_font(caption, size=9, color=BLACK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.first_line_indent = Pt(0)

    if "Manuscript Metadata" not in styles:
        metadata = styles.add_style("Manuscript Metadata", WD_STYLE_TYPE.PARAGRAPH)
    else:
        metadata = styles["Manuscript Metadata"]
    set_style_font(metadata, size=8.5, color=BLACK)
    metadata.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(0)
    metadata.paragraph_format.space_after = Pt(2)
    metadata.paragraph_format.line_spacing = 1.1

    if "English Paper Title" not in styles:
        english_title = styles.add_style("English Paper Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        english_title = styles["English Paper Title"]
    set_style_font(english_title, size=14, color=BLACK, bold=True, east_asia=HEADING_EAST_ASIA_FONT)
    english_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    english_title.paragraph_format.space_before = Pt(0)
    english_title.paragraph_format.space_after = Pt(10)
    english_title.paragraph_format.line_spacing = 1.05

    if "Draft Callout" not in styles:
        callout = styles.add_style("Draft Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Draft Callout"]
    set_style_font(callout, size=8.5, color=BLACK)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(6)
    callout.paragraph_format.line_spacing = 1.15
    callout.paragraph_format.first_line_indent = Pt(0)

    custom_styles = {
        "Author Names": (16, EAST_ASIA_FONT, False, WD_ALIGN_PARAGRAPH.CENTER, 2),
        "Affiliation": (8.5, EAST_ASIA_FONT, False, WD_ALIGN_PARAGRAPH.CENTER, 1),
        "Abstract Heading": (10.5, HEADING_EAST_ASIA_FONT, True, WD_ALIGN_PARAGRAPH.CENTER, 4),
        "Chinese Abstract": (9, EAST_ASIA_FONT, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 3),
        "English Abstract": (10.5, EAST_ASIA_FONT, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 3),
    }
    for style_name, (size, east_asia, bold, alignment, after) in custom_styles.items():
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        set_style_font(style, size=size, color=BLACK, bold=bold, east_asia=east_asia)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.keep_with_next = style_name in {"Author Names", "Abstract Heading"}

    if "Source Code" in styles:
        source_code = styles["Source Code"]
        set_style_font(source_code, size=9, color=BLACK)
        source_code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        source_code.paragraph_format.left_indent = Inches(0.18)
        source_code.paragraph_format.right_indent = Inches(0.12)
        source_code.paragraph_format.space_before = Pt(5)
        source_code.paragraph_format.space_after = Pt(8)
        source_code.paragraph_format.line_spacing = 1.15
        source_code.paragraph_format.first_line_indent = Pt(0)


def remap_heading_levels(doc: Document) -> None:
    title_seen = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        if not title_seen and text.startswith("CAVS："):
            paragraph.style = doc.styles["Title"]
            title_seen = True
            continue
        mapping = {"Heading 2": "Heading 1", "Heading 3": "Heading 2", "Heading 4": "Heading 3"}
        if style_name in mapping:
            paragraph.style = doc.styles[mapping[style_name]]


def style_cover_and_special_paragraphs(doc: Document) -> None:
    image_paragraphs: set[int] = set()
    for shape in doc.inline_shapes:
        parent = shape._inline.getparent()
        while parent is not None and parent.tag != qn("w:p"):
            parent = parent.getparent()
        if parent is not None:
            image_paragraphs.add(id(parent))

    cover_stage = 0
    metadata_index = 0
    abstract_stage: str | None = None
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("CAVS："):
            cover_stage = 1
        elif cover_stage == 1 and text.startswith("CAVS: Contextual"):
            paragraph.style = doc.styles["English Paper Title"]
            for run in paragraph.runs:
                set_run_font(
                    run,
                    size=14,
                    color=BLACK,
                    bold=True,
                    italic=False,
                    east_asia=HEADING_EAST_ASIA_FONT,
                )
            cover_stage = 2
        elif cover_stage == 2 and text and not text.startswith("稿件状态") and text != "摘要":
            metadata_index += 1
            paragraph.style = doc.styles["Author Names" if metadata_index == 1 else "Affiliation"]
            size = 16 if metadata_index == 1 else 8.5
            for run in paragraph.runs:
                set_run_font(run, size=size, color=BLACK)
        if "稿件状态（2026-08-29）" in text:
            paragraph.style = doc.styles["Draft Callout"]
            set_paragraph_shading(paragraph, CALLOUT_FILL)
            set_paragraph_border(paragraph, side="left", color=RULE_GRAY, size=6)
            for run in paragraph.runs:
                set_run_font(run, size=8.5, color=BLACK)
            cover_stage = 3
        if paragraph.style and paragraph.style.name == "Source Code":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_shading(paragraph, CALLOUT_FILL)
            set_paragraph_border(paragraph, side="left", color=RULE_GRAY, size=6)
            for run in paragraph.runs:
                set_run_font(run, size=9, color=BLACK)
        if text == "摘要":
            paragraph.style = doc.styles["Abstract Heading"]
            paragraph.paragraph_format.page_break_before = False
            abstract_stage = "zh"
            cover_stage = 4
        elif text == "Abstract":
            paragraph.style = doc.styles["Abstract Heading"]
            for run in paragraph.runs:
                set_run_font(run, size=10.5, color=BLACK, bold=True)
            abstract_stage = "en"
        elif abstract_stage == "zh" and text:
            paragraph.style = doc.styles["Chinese Abstract"]
        elif abstract_stage == "en" and text:
            paragraph.style = doc.styles["English Abstract"]

        if text.startswith("关键词："):
            paragraph.style = doc.styles["Chinese Abstract"]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(6)
            abstract_stage = None
        elif text.startswith("Keywords:"):
            paragraph.style = doc.styles["English Abstract"]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(8)
            abstract_stage = None

        if id(paragraph._p) in image_paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.keep_with_next = True

        if re.match(r"^(图|表)\s*(?:\d+|B\d+)\b", text):
            paragraph.style = doc.styles["Caption"]
            if text.startswith("表"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9,
                        color=BLACK,
                        bold=True,
                        east_asia=HEADING_EAST_ASIA_FONT,
                    )
            else:
                for run in paragraph.runs:
                    set_run_font(run, size=9, color=BLACK)

        if "【待补" in text or "[to be completed]" in text:
            for run in paragraph.runs:
                if "【待补" in run.text or "[to be completed]" in run.text:
                    run.font.color.rgb = RGBColor.from_string(PENDING_TEXT)
                    run.font.highlight_color = None
            if text.startswith("【待补") or text.startswith("[to be completed]"):
                set_paragraph_shading(paragraph, PENDING_FILL)


def configure_header_footer(doc: Document) -> None:
    for section in doc.sections:
        configure_page(section)
        section.different_first_page_header_footer = True
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        set_paragraph_border(paragraph, side="bottom", color=RULE_GRAY, size=4)
        run = paragraph.add_run("CAVS：严格字节预算下的端云协同目标检测")
        set_run_font(run, size=8, color=BLACK)
        tab = paragraph.add_run("\tV3.3 · 2026-08-29")
        set_run_font(tab, size=8, color=BLACK)
        tabs = paragraph.paragraph_format.tab_stops
        tabs.add_tab_stop(Inches(6.55), alignment=2)

        first_header = section.first_page_header
        first_header_paragraph = first_header.paragraphs[0]
        first_header_paragraph.clear()
        first_header_paragraph.paragraph_format.space_after = Pt(0)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.clear()
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        add_page_field(fp)

        first_footer = section.first_page_footer
        first_fp = first_footer.paragraphs[0]
        first_fp.clear()
        first_fp.paragraph_format.space_before = Pt(0)
        first_fp.paragraph_format.space_after = Pt(0)
        add_page_field(first_fp)


def resize_figures(doc: Document) -> None:
    target_widths = [6.15, 6.15, 6.15, 6.15, 6.15, 6.15]
    alt_texts = [
        "CAVS system overview showing edge detections, action-value prediction, "
        "global or causal scheduling, payload upload, cloud detection, and fusion.",
        "Three-panel line chart of point-estimate AP50–95 versus fixed-byte budget "
        "for CAVS, MORIC, Confidence, and Random on COCO, Oregon, and VOC.",
        "Two-panel VOC action-space ablation showing absolute AP50–95 for CAVS-O, "
        "CAVS-OQ, and CAVS-OQR and the incremental gains from Q90 and ROI.",
        "Two-panel COCO deployment chart showing the cache-to-hardware protocol ladder "
        "and the AP50–95 loss from deployment-aligned offline to formal causal online execution.",
        "Three-panel formal causal-online system chart showing method AP50–95, the CAVS-OQR "
        "end-to-end latency decomposition, and bidirectional bytes per input frame.",
        "Three-panel network-sensitivity chart showing mean latency, P95 latency, and throughput "
        "for CAVS-OQR, CAVS-OQ, ORIC, and MORIC under three network profiles.",
    ]
    for index, shape in enumerate(doc.inline_shapes):
        target = target_widths[index] if index < len(target_widths) else 6.0
        old_width = shape.width
        old_height = shape.height
        shape.width = Inches(target)
        shape.height = int(old_height * shape.width / old_width)
        if index < len(alt_texts):
            doc_pr = shape._inline.docPr
            doc_pr.set("descr", alt_texts[index])
            doc_pr.set("title", f"CAVS manuscript figure {index + 1}")


def display_width(text: str) -> int:
    return sum(2 if ord(char) > 127 else 1 for char in text)


def normalized_widths(table) -> list[int]:
    ncols = len(table.columns)
    header = [cell.text.strip() for cell in table.rows[0].cells]
    joined = "|".join(header)
    if ncols == 6 and "任务价值信号" in joined:
        return [1750, 1700, 1700, 1450, 1300, 1460]
    if ncols == 5 and header and header[0] == "数据集":
        return [1000, 700, 2200, 2800, 2660]
    if ncols == 5 and header[:2] == ["预算", "Original"] and "Q90（及原图比）" in joined:
        # Keep the per-action payload table narrower than the full text block.
        # Its long numeric cells otherwise force LibreOffice to extend the first
        # column into the page margin during PDF verification.
        return [850, 1400, 2450, 2450, 1050]
    if ncols == 7 and header and header[0] == "预算":
        return [900, 1100, 1100, 1100, 1000, 1600, 2560]
    if ncols == 5 and "ROI profile" in joined:
        return [2100, 2100, 1720, 1720, 1720]
    if ncols == 6 and "云动作数" in joined:
        return [900, 1350, 1350, 1600, 2080, 2080]
    if ncols == 8:
        return [900, 700, 900, 1600, 1750, 650, 850, 2010]
    if ncols == 2:
        return [3400, 5960]
    if ncols == 4:
        return [3900, 1820, 1820, 1820]

    maxima = []
    for col_idx in range(ncols):
        max_len = max(display_width(row.cells[col_idx].text.strip()) for row in table.rows)
        maxima.append(min(max(max_len, 8), 28))
    if ncols >= 3:
        maxima[0] *= 1.5
    raw = [max(720, int(CONTENT_WIDTH_DXA * weight / sum(maxima))) for weight in maxima]
    scale = CONTENT_WIDTH_DXA / sum(raw)
    widths = [int(value * scale) for value in raw]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_table_geometry(table, widths: list[int]) -> None:
    tbl = table._tbl
    tblpr = tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(sum(widths)))

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:type"), "dxa")
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        outer_rule = side in {"top", "bottom"}
        edge.set(qn("w:val"), "single" if outer_rule else "nil")
        edge.set(qn("w:sz"), "12" if outer_rule else "0")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), BLACK)

    margins = tblpr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tblpr.append(margins)
    for name, value in CELL_MARGIN_DXA.items():
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(width))


def style_tables(doc: Document) -> None:
    for table in doc.tables:
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = normalized_widths(table)
        set_table_geometry(table, widths)

        header = [cell.text.strip() for cell in table.rows[0].cells]
        is_action_payload_table = (
            len(table.columns) == 5
            and header[:2] == ["预算", "Original"]
            and "Q90（及原图比）" in header
        )

        header_row = table.rows[0]
        trpr = header_row._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        trpr.append(repeat)

        font_size = (
            7.6
            if len(table.columns) >= 7 or is_action_payload_table
            else 8.0
            if len(table.columns) >= 6
            else 8.2
            if len(table.columns) >= 5
            else 8.7
        )
        for row_index, row in enumerate(table.rows):
            row_props = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            cant_split.set(qn("w:val"), "true")
            row_props.append(cant_split)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell, "FFFFFF")
                for side in ("left", "right"):
                    set_cell_border(cell, side=side, val="nil")
                if row_index == 0:
                    set_cell_border(cell, side="bottom", val="single", color=BLACK, size=6)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(1.5)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    # All manuscript tables fit on one A4 page. Chaining their
                    # rows prevents a continuation table from entering the
                    # running-header area in LibreOffice and keeps captions
                    # attached to complete three-line tables in Word.
                    paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if cell == row.cells[0] and row_index > 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            size=font_size,
                            color=BLACK,
                            bold=True if row_index == 0 else None,
                            east_asia=HEADING_EAST_ASIA_FONT if row_index == 0 else EAST_ASIA_FONT,
                        )


def patch_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part.element
    for level in numbering.findall(".//w:lvl", namespaces=numbering.nsmap):
        ppr = level.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            level.append(ppr)
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "279")
        spacing = ppr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            ppr.append(spacing)
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        tabs = ppr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            ppr.append(tabs)
        for child in list(tabs):
            tabs.remove(child)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)

    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.numPr is not None:
            paragraph.paragraph_format.left_indent = Inches(0.375)
            paragraph.paragraph_format.first_line_indent = Inches(-0.194)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.208


def enforce_body_fonts(doc: Document) -> None:
    style_fonts = {
        "Title": (22, HEADING_EAST_ASIA_FONT, False),
        "English Paper Title": (14, HEADING_EAST_ASIA_FONT, True),
        "Author Names": (16, EAST_ASIA_FONT, False),
        "Affiliation": (8.5, EAST_ASIA_FONT, False),
        "Manuscript Metadata": (8.5, EAST_ASIA_FONT, False),
        "Draft Callout": (8.5, EAST_ASIA_FONT, False),
        "Abstract Heading": (10.5, HEADING_EAST_ASIA_FONT, True),
        "Chinese Abstract": (9, EAST_ASIA_FONT, False),
        "English Abstract": (10.5, EAST_ASIA_FONT, False),
        "Heading 1": (14, HEADING_EAST_ASIA_FONT, True),
        "Heading 2": (10.5, HEADING_EAST_ASIA_FONT, True),
        "Heading 3": (10.5, EAST_ASIA_FONT, True),
        "Caption": (9, EAST_ASIA_FONT, False),
        "Source Code": (9, EAST_ASIA_FONT, False),
    }
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else "Normal"
        size, east_asia, force_bold = style_fonts.get(style_name, (10.5, EAST_ASIA_FONT, False))
        if paragraph._p.find(".//m:oMathPara", namespaces=paragraph._p.nsmap) is not None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            set_run_font(
                run,
                size=size,
                color=BLACK,
                east_asia=east_asia,
                bold=True if force_bold else None,
                italic=False if style_name == "English Paper Title" else None,
            )
    for section in doc.sections:
        story_paragraphs = (
            list(section.header.paragraphs)
            + list(section.footer.paragraphs)
            + list(section.first_page_header.paragraphs)
            + list(section.first_page_footer.paragraphs)
        )
        for paragraph in story_paragraphs:
            for run in paragraph.runs:
                if run.text:
                    set_run_font(run, size=8, color=BLACK)


def style_references(doc: Document) -> None:
    in_references = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "参考文献":
            in_references = True
            continue
        if not in_references or not text:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.30)
        paragraph.paragraph_format.first_line_indent = Inches(-0.30)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            set_run_font(run, size=7.5, color=BLACK)


def validate_docx(path: Path) -> dict[str, int]:
    doc = Document(path)
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
        numbering_xml = archive.read("word/numbering.xml")
        core_xml = archive.read("docProps/core.xml")
    stats = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "figures": len(doc.inline_shapes),
        "equations": document_xml.count(b"<m:oMath"),
        "pending_markers": sum("【待补" in paragraph.text for paragraph in doc.paragraphs)
        + sum("【待补" in cell.text for table in doc.tables for row in table.rows for cell in row.cells),
    }
    assert stats["paragraphs"] > 250, stats
    assert stats["tables"] >= 18, stats
    assert stats["figures"] == 6, stats
    assert stats["equations"] >= 10, stats
    for section in doc.sections:
        assert abs(section.page_width.inches - 8.2677) < 0.02
        assert abs(section.page_height.inches - 11.6929) < 0.02
        assert abs(section.left_margin.inches - 0.7874) < 0.02
        assert abs(section.right_margin.inches - 0.7874) < 0.02
    assert b"Times New Roman" in styles_xml
    assert b"Songti SC" in styles_xml
    assert b"Heiti SC" in styles_xml
    assert b"V3.3" in core_xml and b"V3.2" not in core_xml
    assert b'w:left="540"' in numbering_xml and b'w:hanging="279"' in numbering_xml
    for table in doc.tables:
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        assert borders is not None
        assert borders.find(qn("w:top")).get(qn("w:val")) == "single"
        assert borders.find(qn("w:bottom")).get(qn("w:val")) == "single"
        assert borders.find(qn("w:insideV")).get(qn("w:val")) == "nil"
    return stats


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    with tempfile.TemporaryDirectory(prefix="cavs_docx_") as temp_dir:
        temp = Path(temp_dir)
        temp_markdown = temp / "manuscript_word.md"
        raw_docx = temp / "manuscript_raw.docx"
        temp_markdown.write_text(preprocess_markdown(SOURCE.read_text(encoding="utf-8")), encoding="utf-8")
        run_pandoc(temp_markdown, raw_docx)

        doc = Document(raw_docx)
        for section in doc.sections:
            configure_page(section)
        configure_styles(doc)
        remap_heading_levels(doc)
        configure_header_footer(doc)
        resize_figures(doc)
        style_cover_and_special_paragraphs(doc)
        style_tables(doc)
        patch_numbering(doc)
        enforce_body_fonts(doc)
        style_references(doc)

        doc.core_properties.title = "CAVS：严格字节预算下基于上下文图像—动作价值的端云协同目标检测"
        doc.core_properties.subject = "V3.3 中文学术论文排版稿"
        doc.core_properties.keywords = "端云协同, 目标检测, CAVS, 字节预算, RK3588"
        doc.core_properties.comments = "Generated from the canonical V3.3 Markdown manuscript."
        doc.save(OUTPUT)

    stats = validate_docx(OUTPUT)
    print(f"created={OUTPUT}")
    print(" ".join(f"{key}={value}" for key, value in stats.items()))


if __name__ == "__main__":
    build()
