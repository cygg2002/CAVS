#!/usr/bin/env python3
"""Build the JRTIP V0.6 main manuscripts and supplementary Word files."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PACKAGE = ROOT / "paper/manuscript_versions/V0_6_JRTIP_submission"
FIGURES = PACKAGE / "figures"

MAIN_SOURCES = {
    "JRTIP_MANUSCRIPT_V0_6_EN_SUBMISSION.docx": PACKAGE / "JRTIP_MANUSCRIPT_V0_6_EN_SUBMISSION.md",
    "JRTIP_MANUSCRIPT_V0_6_ZH_MASTER.docx": PACKAGE / "JRTIP_MANUSCRIPT_V0_6_ZH_MASTER.md",
}
SUPPLEMENT_SOURCE = PACKAGE / "ONLINE_RESOURCE_1_SUPPLEMENT.md"

FIGURE_FILES_ZH = {
    1: FIGURES / "Fig1_online_edge_cloud_architecture.png",
    2: FIGURES / "Fig2_cafru_offline_training_online_decision.png",
    3: FIGURES / "Fig3_accuracy_upload_compute.png",
    4: FIGURES / "Fig4_weak_network_latency.png",
    5: FIGURES / "Fig5_t41_bounded_queue_deadline.png",
}
FIGURE_FILES_EN = {
    **FIGURE_FILES_ZH,
    1: FIGURES / "Fig1_online_edge_cloud_architecture_en.png",
    2: FIGURES / "Fig2_cafru_offline_training_online_decision_en.png",
    5: FIGURES / "Fig5_t41_bounded_queue_deadline_en.png",
}
FIGURE_CAPTIONS_EN = {
    1: "Online edge-cloud architecture and asynchronous feedback loop",
    2: "Offline utility-label generation and online CAFRU decision",
    3: "COCO deployment-output accuracy-upload-compute trade-off; bubble area represents amortized GFLOPs/image",
    4: "Mean and P95 end-to-end latency under a 10-Mbps uplink and 50-ms RTT",
    5: "T41 main thread, upload worker, and deadline-aware result gate",
}
FIGURE_CAPTIONS_ZH = {
    1: "端云协同在线架构与异步反馈闭环",
    2: "CAFRU 离线效用标签生成与在线决策",
    3: "统一部署输出口径下的 COCO 精度-上传量-计算量折中；气泡面积表示均摊 GFLOPs/image",
    4: "10 Mbps 上行带宽和 50 ms 往返时延下的平均与 P95 端到端延迟",
    5: "T41 主线程、上传工作线程与时限感知结果门控",
}
PLACEHOLDER_HEIGHT_MM = {1: 58, 2: 58, 5: 50}

UNICODE_EQUATIONS = {
    r"R_i=\operatorname{clip}(\operatorname{expand}(b_i,\gamma_i),I_t),": "Rᵢ = clip(expand(bᵢ, γᵢ), Iₜ)",
    r"U(B,Y)=TP(B,Y)-\alpha FP(B,Y),": "U(B,Y) = TP(B,Y) − α FP(B,Y)",
    r"g_f=\frac{U(B_c,Y)-U(B_e,Y)}{\max(1,|Y|)}.": "g_f = [U(B_c,Y) − U(B_e,Y)] / max(1, |Y|)",
    r"g_{r,i}=\frac{U(B_{fr},Y)-U(B_e,Y)}{\max(1,|Y|)}.": "g_r,i = [U(B_fr,Y) − U(B_e,Y)] / max(1, |Y|)",
    r"p_f=\sigma(w_f^\top\tilde x_f+b_f),\qquad\hat g_f=v_f^\top\tilde x_f+a_f,": "p_f = σ(w_fᵀx̃_f + b_f),     ĝ_f = v_fᵀx̃_f + a_f",
    r"p_{r,i}=\sigma(w_r^\top\tilde x_{r,i}+b_r),\qquad\hat g_{r,i}=v_r^\top\tilde x_{r,i}+a_r.": "p_r,i = σ(w_rᵀx̃_r,i + b_r),     ĝ_r,i = v_rᵀx̃_r,i + a_r",
    r"S_f=p_f\max(0,\hat g_f)-\lambda_b C_f,\qquad S_{r,i}=p_{r,i}\max(0,\hat g_{r,i})-\lambda_b C_{r,i},": "S_f = p_f max(0, ĝ_f) − λ_b C_f,     S_r,i = p_r,i max(0, ĝ_r,i) − λ_b C_r,i",
    r"F_{\mathrm{sys}}=F_e+\frac{1}{M}\sum_{t=1}^{M}\left(a_f^tF_c(s_f^t)+\sum_i a_{r,i}^tF_c(s_{r,i}^t)\right),": "F_sys = F_e + (1/M) ∑ₜ [a_fᵗF_c(s_fᵗ) + ∑ᵢ a_r,iᵗF_c(s_r,iᵗ)]",
}


def equation_key(value: str) -> str:
    """Canonicalize a Markdown equation block for editable Word rendering."""
    value = re.sub(r"\\tag\{\d+\}", "", value)
    value = value.replace(r"\qquad", "").replace(r"\quad", "")
    return re.sub(r"\s+", "", value)


UNICODE_EQUATIONS = {equation_key(key): value for key, value in UNICODE_EQUATIONS.items()}


def set_cell_border(cell, **edges) -> None:
    """Apply explicit borders to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, attrs in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in attrs.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row) -> None:
    """Repeat the first row on subsequent pages."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_columns(section, count: int) -> None:
    """Set the section column count and gap."""
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), "340")  # approximately 6 mm
    cols.set(qn("w:equalWidth"), "1")


def configure_section(section, columns: int = 1) -> None:
    """Apply A4 dimensions matching 174-mm full and 84-mm column widths."""
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.header_distance = Mm(7)
    section.footer_distance = Mm(7)
    set_columns(section, columns)


def set_run_font(run, size: float, bold: bool | None = None, italic: bool | None = None, name: str = "Times New Roman") -> None:
    """Set Latin and East Asian fonts consistently."""
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if name == "Times New Roman" else name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_layout(paragraph, size: float = 8.7, first_line: bool = True, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    """Apply compact JRTIP-like body formatting."""
    fmt = paragraph.paragraph_format
    fmt.alignment = alignment
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0.8)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(size + 1.6)
    fmt.first_line_indent = Mm(3.5) if first_line else Mm(0)
    fmt.widow_control = True


def clean_inline(value: str) -> str:
    """Remove Markdown/LaTeX delimiters while retaining readable symbols."""
    value = value.replace("`", "")
    value = value.replace(r"\(", "").replace(r"\)", "")
    value = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", value)
    value = re.sub(r"\\operatorname\{([^{}]+)\}", r"\1", value)
    value = value.replace(r"\{", "{").replace(r"\}", "}")
    value = re.sub(r"_\{([^{}]+)\}", r"_\1", value)
    value = re.sub(r"\^\{([^{}]+)\}", r"^\1", value)
    value = value.replace(r"\mathbb R", "ℝ").replace(r"\mathbb{R}", "ℝ")
    value = value.replace(r"\mathcal A", "A").replace(r"\mathcal{A}", "A")
    value = value.replace(r"\hat g", "ĝ").replace(r"\tilde x", "x̃")
    replacements = {
        r"\ge": "≥",
        r"\alpha": "α",
        r"\lambda": "λ",
        r"\tau": "τ",
        r"\rho": "ρ",
        r"\gamma": "γ",
        r"\sigma": "σ",
        r"\Delta": "Δ",
        r"\top": "ᵀ",
        r"\max": "max",
        r"\in": "∈",
        r"\mathbb 1": "1",
        r"\mathbb{1}": "1",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    value = re.sub(r"\\[A-Za-z]+\{([^{}]+)\}", r"\1", value)
    value = re.sub(r"\\([A-Za-z]+)", r"\1", value)
    inline_replacements = {
        "I_t": "Iₜ",
        "D_e": "Dₑ",
        "D_c": "D꜀",
        "B_e^t": "Bₑᵗ",
        "A_t": "Aₜ",
        "N_t": "Nₜ",
        "b_i": "bᵢ",
        "c_i": "cᵢ",
        "p_i": "pᵢ",
        "x_f": "x_f",
        "x_r,i": "x_r,i",
        "x_1": "x₁",
        "y_1": "y₁",
        "x_2": "x₂",
        "y_2": "y₂",
        "{b_i}_i=1^N_t": "{bᵢ}ᵢ₌₁ᴺᵗ",
    }
    for source, target in inline_replacements.items():
        value = value.replace(source, target)
    return value


def add_inline_runs(paragraph, text: str, size: float = 8.7) -> None:
    """Render simple Markdown emphasis without exposing markup."""
    text = clean_inline(text)
    parts = re.split(r"(\*\*.*?\*\*|\*[^*]+?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size)


def add_page_number(section) -> None:
    """Add a centered PAGE field to the footer."""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_run_font(run, 8)


def setup_document() -> Document:
    """Create a document with stable compact styles."""
    document = Document()
    configure_section(document.sections[0], columns=1)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(8.7)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for section in document.sections:
        add_page_number(section)
    return document


def add_title(document: Document, title: str) -> None:
    """Add the article title."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(title)
    set_run_font(run, 16, bold=True)


def add_author_line(document: Document, text: str, index: int) -> None:
    """Add placeholder author metadata."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(text)
    set_run_font(run, 9.5, italic=index > 0)


def add_heading(document: Document, text: str, level: int) -> None:
    """Add compact decimal headings."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.first_line_indent = Mm(0)
    run = paragraph.add_run(text)
    set_run_font(run, 10.2 if level == 1 else 9.2, bold=True)


def add_body(document: Document, text: str, size: float = 8.7, first_line: bool = True) -> None:
    """Add one justified body paragraph."""
    paragraph = document.add_paragraph()
    set_paragraph_layout(paragraph, size=size, first_line=first_line)
    add_inline_runs(paragraph, " ".join(line.strip() for line in text.splitlines()), size=size)


def add_list_item(document: Document, text: str, ordered: bool) -> None:
    """Add a compact list item."""
    paragraph = document.add_paragraph()
    set_paragraph_layout(paragraph, first_line=False, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.left_indent = Mm(4.5)
    paragraph.paragraph_format.first_line_indent = Mm(-3.5)
    marker_match = re.match(r"^(\d+\.|-)\s+(.*)$", text, flags=re.S)
    marker = marker_match.group(1) if marker_match else ("1." if ordered else "-")
    body = marker_match.group(2) if marker_match else text
    marker_run = paragraph.add_run(f"{marker} ")
    set_run_font(marker_run, 8.7, bold=ordered)
    add_inline_runs(paragraph, body, 8.7)


def add_equation(document: Document, latex: str, number: int) -> None:
    """Add a readable editable linear equation paragraph."""
    normalized = " ".join(line.strip() for line in latex.splitlines())
    equation = UNICODE_EQUATIONS.get(equation_key(normalized), clean_inline(normalized))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"{equation}    ({number})")
    set_run_font(run, 9, name="Cambria Math")


def switch_columns(document: Document, columns: int) -> None:
    """Start a continuous section using the requested columns."""
    section = document.add_section(WD_SECTION.CONTINUOUS)
    configure_section(section, columns=columns)
    section.footer.is_linked_to_previous = True


def add_figure_placeholder(document: Document, number: int) -> None:
    """Reserve realistic layout space for a user-redrawn vector workflow."""
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Mm(170)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row = table.rows[0]
    row.height = Mm(PLACEHOLDER_HEIGHT_MM[number])
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    set_cell_border(
        cell,
        top={"val": "single", "sz": "8", "color": "808080"},
        bottom={"val": "single", "sz": "8", "color": "808080"},
        left={"val": "single", "sz": "8", "color": "808080"},
        right={"val": "single", "sz": "8", "color": "808080"},
    )
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        f"FIGURE {number} PLACEHOLDER\n"
        "Replace with the Visio/draw.io vector figure specified in "
        "FIGURE_DRAWING_SPEC_ZH.md"
    )
    set_run_font(run, 9, bold=True, name="Arial")


def add_figure(document: Document, number: int, caption_text: str, figure_path: Path) -> None:
    """Insert a generated figure or a full-width placeholder."""
    full_width = number in {1, 2, 5}
    if full_width:
        switch_columns(document, 1)
        if figure_path.exists():
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(figure_path), width=Mm(170))
        else:
            add_figure_placeholder(document, number)
    else:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(figure_path), width=Mm(81))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.space_before = Pt(1)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.first_line_indent = Mm(0)
    prefix = caption.add_run(f"Fig. {number} ")
    set_run_font(prefix, 7.7, bold=True)
    body = caption.add_run(caption_text)
    set_run_font(body, 7.7)
    if full_width:
        switch_columns(document, 2)


def strip_cell_markup(value: str) -> str:
    """Clean table-cell Markdown."""
    return clean_inline(value.strip().replace("**", "").replace("<br>", "; "))


def add_table(document: Document, caption_text: str | None, lines: list[str], full_width: bool = True) -> None:
    """Add a compact publication table."""
    keep_together = bool(caption_text and re.search(r"(?:Table|表)\s+5\b", caption_text))
    if full_width:
        switch_columns(document, 1)
    if caption_text:
        caption = document.add_paragraph()
        caption.paragraph_format.keep_with_next = True
        caption.paragraph_format.space_after = Pt(2)
        caption.paragraph_format.first_line_indent = Mm(0)
        add_inline_runs(caption, caption_text.replace("**", ""), 7.8)
    rows = [[strip_cell_markup(cell) for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for index, row in enumerate(rows) if index != 1]
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.allow_autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index > 1 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.keep_with_next = keep_together and row_index < len(rows) - 1
            run = paragraph.add_run(value)
            set_run_font(run, 6.2 if column_count >= 7 else 6.8, bold=row_index == 0)
            border = {"val": "single", "sz": "4", "color": "A6A6A6"}
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E2F3")
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    if full_width:
        switch_columns(document, 2)


def split_blocks(text: str) -> list[str]:
    """Split Markdown into blank-line-delimited blocks."""
    return [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]


def is_table_block(block: str) -> bool:
    """Return whether a block is a Markdown table."""
    lines = block.splitlines()
    return len(lines) >= 2 and lines[0].lstrip().startswith("|") and re.match(r"^\s*\|?[\s:|-]+\|", lines[1]) is not None


def build_main(source: Path, output: Path) -> None:
    """Build a double-column main manuscript from Markdown."""
    blocks = split_blocks(source.read_text(encoding="utf-8"))
    document = setup_document()
    is_chinese = "_ZH_" in source.name
    figure_captions = FIGURE_CAPTIONS_ZH if is_chinese else FIGURE_CAPTIONS_EN
    figure_files = FIGURE_FILES_ZH if is_chinese else FIGURE_FILES_EN
    in_body = False
    equation_number = 0
    pending_table_caption: str | None = None
    author_index = 0

    for block in blocks:
        if block.startswith("# "):
            add_title(document, block[2:].strip())
            continue
        if block.startswith("## Abstract") or block.startswith("## 摘要"):
            heading = "Abstract" if "Abstract" in block else "摘要"
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(heading)
            set_run_font(run, 9, bold=True)
            continue
        if block.startswith("**Keywords:**") or block.startswith("**关键词：**"):
            add_body(document, block, size=8.5, first_line=False)
            continue
        if block.startswith("## "):
            heading = block[3:].strip()
            if re.match(r"^(1 |1\.|1 引言|1 Introduction)", heading) and not in_body:
                switch_columns(document, 2)
                in_body = True
            add_heading(document, heading, 1)
            continue
        if block.startswith("### "):
            add_heading(document, block[4:].strip(), 2)
            continue
        if re.match(r"^\*\*Table \d+|^\*\*表 \d+", block):
            pending_table_caption = block
            continue
        if is_table_block(block):
            add_table(document, pending_table_caption, block.splitlines(), full_width=True)
            pending_table_caption = None
            continue
        figure_match = re.search(r"\[(?:Fig\.|图)\s*(\d+)", block)
        if figure_match and "占位" in block or figure_match and block.startswith("**["):
            number = int(figure_match.group(1))
            add_figure(document, number, figure_captions[number], figure_files[number])
            continue
        if block.startswith(r"\[") and block.endswith(r"\]"):
            equation_number += 1
            add_equation(document, block[2:-2].strip(), equation_number)
            continue
        if re.match(r"^(\d+\.|-)\s+", block):
            add_list_item(document, block, ordered=block[0].isdigit())
            continue
        if not in_body and not block.startswith("##"):
            if author_index < 3:
                add_author_line(document, block, author_index)
                author_index += 1
            else:
                add_body(document, block, size=8.5)
            continue
        add_body(document, block, size=7.6 if block.startswith("[") and re.match(r"^\[\d+\]", block) else 8.7)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def build_supplement(source: Path, output: Path) -> None:
    """Build a single-column supplementary document."""
    blocks = split_blocks(source.read_text(encoding="utf-8"))
    document = setup_document()
    configure_section(document.sections[0], columns=1)
    pending_caption: str | None = None
    author_index = 0
    for block in blocks:
        if block.startswith("# "):
            add_title(document, block[2:].strip())
        elif block.startswith("## "):
            add_heading(document, block[3:].strip(), 1)
        elif block.startswith("### "):
            add_heading(document, block[4:].strip(), 2)
        elif re.match(r"^\*\*Table S\d+|^### Table S\d+", block):
            pending_caption = block.replace("### ", "")
        elif is_table_block(block):
            add_table(document, pending_caption, block.splitlines(), full_width=False)
            pending_caption = None
        elif re.match(r"^(\d+\.|-)\s+", block):
            add_list_item(document, block, ordered=block[0].isdigit())
        elif author_index < 2 and not block.startswith("#"):
            add_author_line(document, block, author_index)
            author_index += 1
        else:
            add_body(document, block)
    document.save(output)


def main() -> None:
    """Build all V0.6 Word artifacts."""
    for output_name, source in MAIN_SOURCES.items():
        build_main(source, PACKAGE / output_name)
        print(f"Built {output_name}")
    build_supplement(SUPPLEMENT_SOURCE, PACKAGE / "ONLINE_RESOURCE_1_SUPPLEMENT.docx")
    print("Built ONLINE_RESOURCE_1_SUPPLEMENT.docx")


if __name__ == "__main__":
    main()
