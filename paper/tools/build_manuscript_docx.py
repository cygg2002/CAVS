#!/usr/bin/env python3
"""Build a formatted DOCX manuscript from the current JRTIP Markdown draft."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ZH_FONT = "宋体"
ZH_HEADING_FONT = "黑体"
EN_FONT = "Times New Roman"
MONO_FONT = "Courier New"


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None,
                 color: str | None = None, east_asia: str = ZH_FONT, latin: str = EN_FONT) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, first_line: bool = False, space_after: float = 6,
                         align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.15
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align


def shade_element(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    element.get_or_add_tcPr().append(shd)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, 9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = EN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), ZH_FONT)
    normal.font.size = Pt(10.5)

    for name, size in (("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 11), ("Heading 4", 10.5)):
        style = styles[name]
        style.font.name = EN_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), ZH_HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = EN_FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), ZH_FONT)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)


def add_inline_markdown(paragraph, text: str, size: float = 10.5, first_run_bold: bool = False) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    first = True
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size, bold=(first and first_run_bold))
            first = False
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size, east_asia=MONO_FONT, latin=MONO_FONT)
        first = False
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size, bold=(first and first_run_bold))


def split_md_table_row(line: str) -> list[str]:
    row = line.strip().strip("|")
    return [cell.strip() for cell in row.split("|")]


def is_md_separator(line: str) -> bool:
    if not line.strip().startswith("|"):
        return False
    cells = split_md_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows = [split_md_table_row(line) for idx, line in enumerate(table_lines) if idx != 1]
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    small_font = 7.2 if col_count >= 10 else 8.2 if col_count >= 7 else 9.0
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                shade_element(cell._tc, "D9EAF7")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            text = row[c_idx] if c_idx < len(row) else ""
            add_inline_markdown(paragraph, text, small_font, first_run_bold=(r_idx == 0))
    doc.add_paragraph()


def add_code_block(doc: Document, code: list[str]) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, first_line=False, space_after=6)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    run = paragraph.add_run("\n".join(code))
    set_run_font(run, 8.5, east_asia=MONO_FONT, latin=MONO_FONT)


def add_image(doc: Document, base_dir: Path, alt: str, ref: str) -> None:
    image_path = base_dir / ref
    if not image_path.exists():
        paragraph = doc.add_paragraph()
        add_inline_markdown(paragraph, f"[Missing figure: {ref}] {alt}", 9)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(15.7))
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_markdown(caption, alt, 9)


def add_title_block(doc: Document, lines: list[str]) -> int:
    title = lines[0].lstrip("#").strip()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(6)
    run = title_p.add_run(title)
    set_run_font(run, 16, bold=True, east_asia=ZH_HEADING_FONT, latin=EN_FONT)

    idx = 1
    metadata: list[str] = []
    while idx < len(lines) and not lines[idx].startswith("## "):
        if lines[idx].strip():
            metadata.append(lines[idx].strip())
        idx += 1

    for raw in metadata:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        size = 10 if any(raw.startswith(prefix) for prefix in ("目标期刊", "稿件版本", "更新时间", "写作状态")) else 10.5
        add_inline_markdown(paragraph, raw, size)
    doc.add_paragraph()
    return idx


def build_docx(source_md: Path, output_docx: Path) -> dict:
    base_dir = source_md.parent
    lines = source_md.read_text(encoding="utf-8").splitlines()

    doc = Document()
    configure_document(doc)
    i = add_title_block(doc, lines)

    image_count = 0
    table_count = 0
    code_count = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            add_code_block(doc, code)
            code_count += 1
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_md_separator(lines[i + 1]):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            table_count += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            add_image(doc, base_dir, image_match.group(1), image_match.group(2))
            image_count += 1
            i += 1
            continue

        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            text = stripped[level:].strip()
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline_markdown(paragraph, text, 13 if level == 1 else 12 if level == 2 else 11)
            i += 1
            continue

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            set_paragraph_format(paragraph, first_line=False, space_after=6)
            paragraph.paragraph_format.left_indent = Cm(0.5)
            add_inline_markdown(paragraph, stripped.lstrip("> ").strip(), 9.5)
            i += 1
            continue

        if re.fullmatch(r"\*\*表\s*[^*]+?\*\*", stripped):
            paragraph = doc.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(paragraph, stripped.strip("*"), 9, first_run_bold=True)
            i += 1
            continue

        paragraph = doc.add_paragraph()
        set_paragraph_format(paragraph, first_line=True, space_after=5)
        if re.match(r"^\d+\.\s+", stripped) or stripped.startswith("- "):
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0.5)
        add_inline_markdown(paragraph, stripped, 10.5)
        i += 1

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)
    return {
        "source_md": str(source_md),
        "output_docx": str(output_docx),
        "images": image_count,
        "tables": table_count,
        "code_blocks": code_count,
        "lines": len(lines),
    }


def make_submission_markdown(source_md: Path, output_md: Path) -> Path:
    lines = source_md.read_text(encoding="utf-8").splitlines()
    cleaned: list[str] = []
    skip_section = False
    skip_heading = ""
    for line in lines:
        stripped = line.strip()
        if stripped == "## 图表清单":
            skip_section = True
            skip_heading = stripped
            continue
        if stripped == "## 投稿版压缩与后续补充建议":
            skip_section = True
            skip_heading = stripped
            continue
        if skip_section and stripped.startswith("## "):
            skip_section = False
            skip_heading = ""
        if skip_section:
            continue
        if any(stripped.startswith(prefix) for prefix in ("目标期刊：", "稿件版本：", "更新时间：", "写作状态：")):
            continue
        cleaned.append(line)
    output_md.parent.mkdir(parents=True, exist_ok=True)
    output_md.write_text("\n".join(cleaned).strip() + "\n", encoding="utf-8")
    return output_md


def copy_package_files(source_md: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_md, output_dir / "manuscript_source.md")
    revision_notes = source_md.parent / "V0_5_REVISION_NOTES.md"
    if revision_notes.exists():
        shutil.copy2(revision_notes, output_dir / revision_notes.name)
    source_figures = source_md.parent / "figures"
    if source_figures.exists():
        dest_figures = output_dir / "figures"
        dest_figures.mkdir(exist_ok=True)
        for fig in sorted(source_figures.glob("*")):
            if fig.is_file() and not fig.name.startswith("~$"):
                shutil.copy2(fig, dest_figures / fig.name)


def write_readme(output_dir: Path, stats: dict) -> None:
    readme = f"""# V0.5 Standard Word Package

This folder contains a formatted Word manuscript generated from the latest V0.5 revised Markdown draft.

- Source Markdown copy: `manuscript_source.md`
- Word manuscript: `{Path(stats["output_docx"]).name}`
- Embedded figures: {stats["images"]}
- Markdown tables converted to Word tables: {stats["tables"]}
- Code / algorithm blocks: {stats["code_blocks"]}
- Source lines: {stats["lines"]}

Formatting notes:
- A4 page, academic margins, page number footer.
- Chinese body font uses Songti; headings use Heiti; English text uses Times New Roman.
- Figures are centered and captioned from Markdown alt text.
- Wide tables use smaller font to keep the manuscript complete in portrait pages.

Remaining submission tasks:
- Fill author, affiliation, correspondence, funding, conflict-of-interest, code availability, and author contributions.
- Decide whether SpeciesNet-generated annotations should be described as automatic reference annotations or pseudo-GT in the English submission.
- Replace PNG figures with vector PDF/SVG exports if required by the target journal.
"""
    (output_dir / "README.md").write_text(readme, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-md", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--docx-name", default="JRTIP_MANUSCRIPT_V0_5_STANDARD_FULL_ZH.docx")
    parser.add_argument("--submission-clean", action="store_true")
    args = parser.parse_args()

    source_md = args.source_md.resolve()
    output_dir = args.output_dir.resolve()
    copy_package_files(source_md, output_dir)
    build_source_md = source_md
    if args.submission_clean:
        build_source_md = make_submission_markdown(source_md, output_dir / "manuscript_submission_clean.md")
    docx_path = output_dir / args.docx_name
    stats = build_docx(build_source_md, docx_path)
    stats["original_source_md"] = str(source_md)
    stats["submission_clean"] = bool(args.submission_clean)
    (output_dir / "build_manifest.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_readme(output_dir, stats)
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
